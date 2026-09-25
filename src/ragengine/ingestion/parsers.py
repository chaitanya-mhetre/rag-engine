"""Format-specific parsers behind one `Parser` protocol.

Each parser turns raw bytes into `Section`s that remember their heading path (and page, for
PDFs). Structure is kept because the recursive chunker and the citations both use it.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Protocol

from ragengine.ingestion.cleaning import clean_text
from ragengine.models import ParsedDocument, Section

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class UnsupportedFormatError(ValueError):
    """Raised when no parser handles the detected mime type."""


class Parser(Protocol):
    mime_types: tuple[str, ...]

    def parse(self, data: bytes, filename: str) -> ParsedDocument: ...


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _title_from_filename(filename: str) -> str:
    return Path(filename).stem.replace("_", " ").replace("-", " ").strip() or filename


class _HeadingStack:
    """Tracks the current heading path. A level-N heading closes every open heading >= N."""

    def __init__(self) -> None:
        self._stack: list[tuple[int, str]] = []

    def push(self, level: int, text: str) -> None:
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, text))

    @property
    def path(self) -> tuple[str, ...]:
        return tuple(text for _, text in self._stack)


class TextParser:
    mime_types: tuple[str, ...] = ("text/plain",)

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        text = clean_text(_decode(data))
        paragraphs = [p for p in text.split("\n\n") if p.strip()]
        return ParsedDocument(
            title=_title_from_filename(filename),
            mime_type="text/plain",
            sections=[Section(p) for p in paragraphs],
        )


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")


class MarkdownParser:
    """Tracks ATX headings (`#`, `##`, ...) so each section knows its heading path.

    A line-based parser is enough for headings; fenced code blocks are skipped for heading
    detection so a `# comment` inside code is not treated as a heading.
    """

    mime_types: tuple[str, ...] = ("text/markdown",)

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        text = _decode(data)
        sections: list[Section] = []
        headings = _HeadingStack()
        buffer: list[str] = []
        in_fence = False
        title: str | None = None

        def flush() -> None:
            body = clean_text("\n".join(buffer))
            for para in (p for p in body.split("\n\n") if p.strip()):
                sections.append(Section(para, headings.path))
            buffer.clear()

        for line in text.splitlines():
            if line.strip().startswith(("```", "~~~")):
                in_fence = not in_fence
                buffer.append(line)
                continue
            match = None if in_fence else _MD_HEADING.match(line)
            if match:
                flush()
                level, heading = len(match.group(1)), match.group(2).strip()
                headings.push(level, heading)
                if title is None and level == 1:
                    title = heading
            else:
                buffer.append(line)
        flush()
        return ParsedDocument(
            title=title or _title_from_filename(filename),
            mime_type="text/markdown",
            sections=sections,
        )


class HtmlParser:
    mime_types: tuple[str, ...] = ("text/html",)
    _BLOCK_TAGS = ("p", "li", "pre", "blockquote", "td", "th", "dd", "dt")

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        from bs4 import BeautifulSoup, Tag

        soup = BeautifulSoup(_decode(data), "html.parser")
        for junk in soup(["script", "style", "nav", "footer", "noscript"]):
            junk.decompose()
        title_tag = soup.find("title")
        title = title_tag.get_text(strip=True) if title_tag else _title_from_filename(filename)

        sections: list[Section] = []
        stack = _HeadingStack()
        headings = {f"h{i}" for i in range(1, 7)}
        root = soup.body or soup
        for el in root.find_all(list(headings) + list(self._BLOCK_TAGS)):
            if not isinstance(el, Tag):
                continue
            text = clean_text(el.get_text(" ", strip=True))
            if not text:
                continue
            if el.name in headings:
                stack.push(int(el.name[1]), text)
                continue
            # Skip blocks nested in another block tag we already emit (e.g. <p> inside <li>).
            if el.find_parent(self._BLOCK_TAGS) is not None:
                continue
            sections.append(Section(text, stack.path))
        return ParsedDocument(title=title, mime_type="text/html", sections=sections)


class PdfParser:
    """Text-layer PDFs only; scanned PDFs (no text layer) come back empty. OCR is out of scope."""

    mime_types: tuple[str, ...] = ("application/pdf",)

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedFormatError("encrypted PDFs are not supported")
        sections: list[Section] = []
        for page_no, page in enumerate(reader.pages, start=1):
            text = clean_text(page.extract_text() or "")
            for para in (p for p in text.split("\n\n") if p.strip()):
                sections.append(Section(para, (), page_no))
        meta_title = reader.metadata.title if reader.metadata else None
        return ParsedDocument(
            title=meta_title or _title_from_filename(filename),
            mime_type="application/pdf",
            sections=sections,
            metadata={"pages": len(reader.pages)},
        )


class DocxParser:
    """Uses Word's built-in "Heading N" styles to build the heading path."""

    mime_types: tuple[str, ...] = (DOCX_MIME,)

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        import docx

        document = docx.Document(io.BytesIO(data))
        sections: list[Section] = []
        stack = _HeadingStack()
        title: str | None = document.core_properties.title or None
        for para in document.paragraphs:
            text = clean_text(para.text)
            if not text:
                continue
            style = para.style.name if para.style is not None and para.style.name else ""
            if style == "Title":
                title = title or text
                continue
            if style.startswith("Heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:
                    level = 1
                stack.push(level, text)
                continue
            sections.append(Section(text, stack.path))
        return ParsedDocument(
            title=title or _title_from_filename(filename),
            mime_type=self.mime_types[0],
            sections=sections,
        )


_EXTENSION_MIME = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": DocxParser.mime_types[0],
}


def detect_mime(data: bytes, filename: str) -> str:
    """Sniff magic bytes first; fall back to the extension only for text formats.

    Trusting the extension alone would let `evil.pdf` that is really a zip reach the PDF parser.
    """
    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith(b"PK\x03\x04"):
        if b"word/" in data[:4096] or filename.lower().endswith(".docx"):
            return DocxParser.mime_types[0]
        raise UnsupportedFormatError("zip archive that is not a DOCX document")
    ext = Path(filename).suffix.lower()
    mime = _EXTENSION_MIME.get(ext)
    if mime in ("application/pdf", DocxParser.mime_types[0]):
        raise UnsupportedFormatError(f"{filename}: extension says {ext} but content does not match")
    head = data[:512].lstrip().lower()
    if mime is None and (head.startswith(b"<!doctype html") or head.startswith(b"<html")):
        return "text/html"
    if mime is None:
        raise UnsupportedFormatError(f"unsupported file type: {filename}")
    return mime


class ParserRegistry:
    def __init__(self, parsers: list[Parser] | None = None) -> None:
        self._by_mime: dict[str, Parser] = {}
        for parser in parsers or [
            TextParser(),
            MarkdownParser(),
            HtmlParser(),
            PdfParser(),
            DocxParser(),
        ]:
            for mime in parser.mime_types:
                self._by_mime[mime] = parser

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        mime = detect_mime(data, filename)
        parser = self._by_mime.get(mime)
        if parser is None:
            raise UnsupportedFormatError(mime)
        return parser.parse(data, filename)
